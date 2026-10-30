#!/usr/bin/env python3
"""Import the official EDBO 2026 C7 workbook into the site snapshot format.

Usage:
    python scripts/import-edbo-official-2026.py path/to/source.xlsx

The workbook is the official response from the State Enterprise
"Inforesurs". It contains one row per competitive offer. The importer keeps
all applications as the default view, preserves a separate 1st/2nd-priority
view, combines study forms and offer types, and calculates the main average
score weighted by the number of applications admitted to the competition
(not by all submitted applications).
"""

from __future__ import annotations

import json
import re
import sys
from collections import defaultdict
from datetime import datetime, timezone, timedelta
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - friendly CLI failure
    raise SystemExit("Install openpyxl first: python -m pip install openpyxl") from exc


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "data" / "2026-current.json"
MIN_APPLICATIONS = {"bachelor": 20, "master": 15}

UID_TO_SLUG = {
    41: "knu", 282: "lnu", 79: "naukma", 244: "ucu", 28: "onu",
    62: "karazin", 111: "dnu", 73: "znu", 44: "vnu", 207: "uzhnu",
    101: "cnu", 6945: "kubg", 246: "donnu", 198: "cpu",
    81: "lnu-shev", 19: "mdu", 61: "chnu", 168: "sumdu",
    6704: "npu", 337: "kneu",
}
NAME_OVERRIDE = {
    337: "Київський національний економічний університет імені Вадима Гетьмана",
    6594: "Державний торговельно-економічний університет (Київ)",
}
PROGRAM_EXCEPTIONS = {
    (41, "Цифрові медіа"),
    (6945, "Контент-продюсування цифрових медіапроєктів"),
    (6945, "Міжнародні медіа та цифрові комунікації"),
}
SLUG_SHORT_OVERRIDE = {"kneu": "КНЕУ"}
PALETTE_HUES = [350, 205, 268, 140, 24, 12, 60, 90, 320, 200, 150, 45, 300, 260, 18, 100, 220, 280, 170, 330, 210, 30, 240, 60, 130]
SLUG_HUE = {
    "knu": 350, "lnu": 205, "naukma": 268, "ucu": 140, "onu": 24,
    "karazin": 12, "dnu": 60, "znu": 90, "vnu": 200, "uzhnu": 150,
    "cnu": 45, "kubg": 300, "donnu": 260, "cpu": 18, "lnu-shev": 100,
    "mdu": 220, "chnu": 280, "sumdu": 170, "npu": 330, "kneu": 130,
}
STOPWORDS = {
    "національний", "державний", "приватний", "вищий", "вищої", "заклад",
    "закладу", "освіти", "освіта", "університет", "університету", "інститут",
    "інституту", "академія", "академії", "коледж", "навчальний", "навчальної",
    "імені", "ім", "та", "і", "в", "у", "з", "до", "на", "як", "або",
    "the", "of", "for",
}

EXPECTED_HEADERS = [
    "Код закладу освіти в ЄДЕБО",
    "Назва закладу освіти",
    "Освітні програми",
    "ID конкурсної пропозиції",
    "Освітній ступінь",
    "Форма навчання",
    "Вид конкурсної пропозиції",
    "Усього подано заяв",
    "Кількість заяв, допущених до конкурсу",
    "Середній конкурсний бал заяв, допущених до конкурсу",
    "Кількість заяв першого та другого пріоритетів",
    "Середній конкурсний бал заяв першого та другого пріоритетів",
    "Кількість зарахованих на бюджет",
    "Кількість зарахованих на контракт",
]


def number(value, default=0):
    if value in (None, ""):
        return default
    return float(value)


def short_name(name: str) -> str:
    cleaned = re.sub(r"[\"'«»“”‘’()]", " ", name)
    words = [w for w in re.split(r"\s+", cleaned) if len(w) > 1 and w.lower() not in STOPWORDS]
    source = words or [w for w in re.split(r"\s+", cleaned) if w]
    parts = source if len(source) >= 3 else [p for w in source for p in w.split("-")]
    letters = "".join(w[0] for w in parts[:4] if w).upper()
    return letters or name[:4].upper()


def is_journalism(uid: int, program: str) -> bool:
    return bool(re.search("журналіст", program, re.IGNORECASE)) or (uid, program) in PROGRAM_EXCEPTIONS


def add_offer(rec: dict, row: tuple) -> None:
    uid, name, program, _offer_id, _degree, _form, offer_type, submitted, admitted, avg_score, p12, p12_score, enrolled_budget, enrolled_contract = row
    submitted = int(number(submitted))
    admitted = int(number(admitted))
    avg_score = number(avg_score)
    p12 = int(number(p12))
    p12_score = number(p12_score)
    is_budget = offer_type != "Небюджетна"

    rec["name"] = str(name).strip()
    rec["programs"].add(str(program).strip())
    rec["applicationsTotal"] += submitted
    rec["admitted"] += admitted
    rec["weightedScore"] += avg_score * admitted
    rec["applicationsP12Total"] += p12
    rec["weightedScoreP12"] += p12_score * p12
    rec["enrolledBudget"] += int(number(enrolled_budget))
    rec["enrolledContract"] += int(number(enrolled_contract))
    if is_budget:
        rec["applicationsOpenTotal"] += submitted
        rec["admittedOpen"] += admitted
        rec["weightedScoreOpen"] += avg_score * admitted
        rec["applicationsOpenP12Total"] += p12
        rec["weightedScoreOpenP12"] += p12_score * p12


def build_rows(grouped: dict, level: str) -> list[dict]:
    result = []
    for uid, rec in grouped.items():
        program_count = len(rec["programs"])
        if rec["applicationsTotal"] < MIN_APPLICATIONS[level] or rec["admitted"] <= 0 or program_count <= 0:
            continue
        slug = UID_TO_SLUG.get(uid)
        score = round(rec["weightedScore"] / rec["admitted"], 1)
        applications = round(rec["applicationsTotal"] / program_count, 1)
        applications_open = round(rec["applicationsOpenTotal"] / program_count, 1)
        score_open = round(rec["weightedScoreOpen"] / rec["admittedOpen"], 1) if rec["admittedOpen"] else None
        applications_p12 = round(rec["applicationsP12Total"] / program_count, 1)
        score_p12 = round(rec["weightedScoreP12"] / rec["applicationsP12Total"], 1) if rec["applicationsP12Total"] else None
        applications_open_p12 = round(rec["applicationsOpenP12Total"] / program_count, 1)
        score_open_p12 = round(rec["weightedScoreOpenP12"] / rec["applicationsOpenP12Total"], 1) if rec["applicationsOpenP12Total"] else None
        result.append({
            "id": slug or f"edbo{uid}",
            "name": NAME_OVERRIDE.get(uid, rec["name"]),
            "short": SLUG_SHORT_OVERRIDE.get(slug, short_name(rec["name"])),
            "hue": SLUG_HUE.get(slug, PALETTE_HUES[uid % len(PALETTE_HUES)]),
            "score": score,
            "applications": applications,
            "applicationsTotal": rec["applicationsTotal"],
            "programCount": program_count,
            "admitted": rec["admitted"],
            "applicationsP12Total": rec["applicationsP12Total"],
            "applicationsP12": applications_p12,
            "scoreP12": score_p12,
            "applicationsOpenTotal": rec["applicationsOpenTotal"],
            "applicationsOpen": applications_open,
            "admittedOpen": rec["admittedOpen"],
            "scoreOpen": score_open,
            "applicationsOpenP12Total": rec["applicationsOpenP12Total"],
            "applicationsOpenP12": applications_open_p12,
            "scoreOpenP12": score_open_p12,
            "enrolledBudget": rec["enrolledBudget"],
            "enrolledContract": rec["enrolledContract"],
        })
    result.sort(key=lambda row: (row["score"], row["applicationsTotal"]), reverse=True)
    for rank, row in enumerate(result, 1):
        row["rank"] = rank
    return result


def empty_record() -> dict:
    return {
        "name": "", "programs": set(), "applicationsTotal": 0, "admitted": 0,
        "weightedScore": 0.0, "applicationsOpenTotal": 0, "admittedOpen": 0,
        "weightedScoreOpen": 0.0, "applicationsP12Total": 0, "weightedScoreP12": 0.0,
        "applicationsOpenP12Total": 0, "weightedScoreOpenP12": 0.0,
        "enrolledBudget": 0, "enrolledContract": 0,
    }


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Usage: python scripts/import-edbo-official-2026.py path/to/source.xlsx")
    source = Path(sys.argv[1]).expanduser().resolve()
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(v).strip() if v is not None else "" for v in next(rows)]
    if headers != EXPECTED_HEADERS:
        raise SystemExit("Unexpected workbook columns; the official EDBO export format may have changed")

    grouped = {"bachelor": defaultdict(empty_record), "master": defaultdict(empty_record)}
    source_rows = 0
    included_rows = 0
    for row in rows:
        source_rows += 1
        uid = int(number(row[0]))
        program = str(row[2] or "").strip()
        degree = str(row[4] or "").strip().lower()
        if not is_journalism(uid, program):
            continue
        level = "bachelor" if "бакалавр" in degree else "master" if "магістр" in degree else None
        if not level:
            continue
        add_offer(grouped[level][uid], row)
        included_rows += 1

    bachelor = build_rows(grouped["bachelor"], "bachelor")
    master = build_rows(grouped["master"], "master")
    modified = workbook.properties.modified or datetime(2026, 10, 27, 14, 10, 16)
    if modified.tzinfo is None:
        modified = modified.replace(tzinfo=timezone(timedelta(hours=2)))
    snapshot = {
        "date": modified.date().isoformat(),
        "asOf": modified.isoformat(timespec="seconds"),
        "final": True,
        "preliminary": False,
        "source": {
            "kind": "official-edbo-xlsx",
            "file": source.name,
            "sheet": sheet.title,
            "rows": source_rows,
            "includedJournalismOffers": included_rows,
            "received": "2026-10-28",
        },
        "bachelor": bachelor,
        "master": master,
        "totalApplications": {
            "bachelor": sum(r["applicationsTotal"] for r in bachelor),
            "master": sum(r["applicationsTotal"] for r in master),
        },
        "totalAdmitted": {
            "bachelor": sum(r["admitted"] for r in bachelor),
            "master": sum(r["admitted"] for r in master),
        },
        "totalApplicationsOpen": {
            "bachelor": sum(r["applicationsOpenTotal"] for r in bachelor),
            "master": sum(r["applicationsOpenTotal"] for r in master),
        },
        "totalApplicationsP12": {
            "bachelor": sum(r["applicationsP12Total"] for r in bachelor),
            "master": sum(r["applicationsP12Total"] for r in master),
        },
        "totalApplicationsOpenP12": {
            "bachelor": sum(r["applicationsOpenP12Total"] for r in bachelor),
            "master": sum(r["applicationsOpenP12Total"] for r in master),
        },
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Imported {included_rows} journalism offer rows from {source_rows}; wrote {OUTPUT}")
    print(f"Bachelor: {len(bachelor)} universities, master: {len(master)} universities")


if __name__ == "__main__":
    main()
